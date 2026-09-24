"""
One path for every patient upload, whether from the Documents page or Ask Aira:
validate -> read the text -> keep the original, encrypted (documents/storage.py).
"""
from __future__ import annotations

import io
import logging

from . import storage
from .models import UploadedDocument

logger = logging.getLogger(__name__)

MAX_UPLOAD_BYTES = 15 * 1024 * 1024  # 15 MB
ALLOWED = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".heic": "image/heic",
}
ALLOWED_EXTENSIONS = set(ALLOWED)


class UploadRejected(Exception):
    pass


def extension(name: str) -> str:
    name = (name or "").lower().strip()
    return "." + name.rsplit(".", 1)[-1] if "." in name else ""


def _sniff_ok(ext: str, data: bytes) -> bool:
    """Make sure the bytes match the extension, so renamed files can't sneak through."""
    head = data[:16]
    if ext == ".pdf":
        return head.startswith(b"%PDF")
    if ext == ".png":
        return head.startswith(b"\x89PNG")
    if ext in (".jpg", ".jpeg"):
        return head.startswith(b"\xff\xd8")
    if ext == ".webp":
        return head[:4] == b"RIFF" and head[8:12] == b"WEBP"
    if ext == ".heic":
        return head[4:8] == b"ftyp"
    if ext == ".docx":
        return head.startswith(b"PK")
    return True  # .txt


def extract_text(name: str, data: bytes) -> str:
    name = (name or "").lower()
    try:
        if name.endswith(".pdf"):
            import fitz

            doc = fitz.open(stream=data, filetype="pdf")
            try:
                return "\n".join(page.get_text() for page in doc).strip()
            finally:
                doc.close()
        if name.endswith(".docx"):
            import docx

            return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs).strip()
        if name.endswith((".png", ".jpg", ".jpeg", ".webp")):
            try:
                import pytesseract
                from PIL import Image

                return (pytesseract.image_to_string(Image.open(io.BytesIO(data))) or "").strip()
            except Exception as exc:
                logger.info("Image OCR unavailable or failed: %s", exc)
                return ""
        if name.endswith(".txt"):
            return data.decode("utf-8", errors="ignore").strip()
    except Exception:
        logger.exception("Document text extraction failed for %s", name)
    return ""


def check(uploaded) -> None:
    """Validate without saving, so a batch can be rejected before anything is stored."""
    name = (getattr(uploaded, "name", "") or "file").rsplit("/", 1)[-1][:255]
    ext = extension(name)
    if ext not in ALLOWED:
        raise UploadRejected(f"“{name}” isn’t a type Aira can read. Use a photo, PDF, Word file, or text file.")
    if uploaded.size and uploaded.size > MAX_UPLOAD_BYTES:
        raise UploadRejected(f"“{name}” is over 15 MB. Try a smaller photo or a shorter PDF.")
    head = uploaded.read(16)
    uploaded.seek(0)
    if not head:
        raise UploadRejected(f"“{name}” looks empty.")
    if not _sniff_ok(ext, head):
        raise UploadRejected(f"“{name}” doesn’t look like a real {ext[1:].upper()} file.")


def image_data_url(data: bytes, max_side: int = 1600) -> str | None:
    """A downsized JPEG data URL for the vision model, or None if the image can't be read."""
    import base64

    try:
        from PIL import Image, ImageOps

        img = ImageOps.exif_transpose(Image.open(io.BytesIO(data)))
        img = img.convert("RGB")
        img.thumbnail((max_side, max_side))
        out = io.BytesIO()
        img.save(out, format="JPEG", quality=85)
        return "data:image/jpeg;base64," + base64.b64encode(out.getvalue()).decode()
    except Exception as exc:
        logger.info("Could not prepare image for the model: %s", exc)
        return None


def ingest(user, uploaded, *, kind=UploadedDocument.Kind.OTHER, source="", chat_message=None) -> UploadedDocument:
    """Validate and save one upload. Raises UploadRejected with a patient-friendly message."""
    name = (getattr(uploaded, "name", "") or "file").rsplit("/", 1)[-1][:255]
    ext = extension(name)
    if ext not in ALLOWED:
        raise UploadRejected(f"“{name}” isn’t a type Aira can read. Use a photo, PDF, Word file, or text file.")
    if uploaded.size and uploaded.size > MAX_UPLOAD_BYTES:
        raise UploadRejected(f"“{name}” is over 15 MB. Try a smaller photo or a shorter PDF.")
    data = uploaded.read()
    if not data:
        raise UploadRejected(f"“{name}” looks empty.")
    if not _sniff_ok(ext, data):
        raise UploadRejected(f"“{name}” doesn’t look like a real {ext[1:].upper()} file.")

    if kind not in UploadedDocument.Kind.values:
        kind = UploadedDocument.Kind.OTHER

    doc = UploadedDocument(
        user=user,
        kind=kind,
        source_institution=(source or "").strip()[:255],
        extracted_text=extract_text(name, data),
        original_name=name,
        content_type=ALLOWED[ext],
        size_bytes=len(data),
        chat_message=chat_message,
    )
    if storage.is_enabled():
        try:
            doc.file_key = storage.save(data)
        except Exception:
            # Keep the text even if the file store is down; the patient isn't blocked.
            logger.exception("Storing the original file failed; keeping extracted text only")
    doc.save()
    doc._raw_bytes = data  # for the current request only (e.g. to show a photo to the model)
    return doc
